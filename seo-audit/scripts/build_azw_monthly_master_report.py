import sys
from datetime import date
from pathlib import Path

import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"C:\Users\yasir\Documents\Final-Capstone\seo-audit")
OUTDIR = ROOT / "reports" / "AZWebCorp-Google-Reports-2026-08-30"
DOCX = OUTDIR / "AZ-Web-Corp-Monthly-Comparison-Master-Report-BRANDED.docx"
XLSX = OUTDIR / "AZ-Web-Corp-Monthly-Comparison-Data.xlsx"
sys.path.insert(0, str(ROOT))
from gsc_query_oauth import query as gsc_query
from scripts.recover_azw_ga4_from_wrong_property import run as recover_ga4
from scripts.document_branding import apply_branding

CURRENT_START, CURRENT_END = date(2026, 8, 1), date(2026, 8, 27)
PREVIOUS_START, PREVIOUS_END = date(2026, 7, 1), date(2026, 7, 27)


def gsc(dimensions=None, start=CURRENT_START, end=CURRENT_END, search_type="web"):
    body = {"startDate": start.isoformat(), "endDate": end.isoformat(), "type": search_type, "rowLimit": 25000, "dataState": "final"}
    if dimensions:
        body["dimensions"] = dimensions
    raw = gsc_query("sc-domain:azwebcorp.com", body)
    rows = []
    for item in raw.get("rows", []):
        row = {d: v for d, v in zip(dimensions or [], item.get("keys", []))}
        row.update({k: item.get(k, 0) for k in ("clicks", "impressions", "ctr", "position")})
        rows.append(row)
    return pd.DataFrame(rows)


def merge_months(current, previous, key):
    cols = [key, "clicks", "impressions", "ctr", "position"]
    c = current.reindex(columns=cols).rename(columns={x: f"current_{x}" for x in cols if x != key})
    p = previous.reindex(columns=cols).rename(columns={x: f"previous_{x}" for x in cols if x != key})
    m = c.merge(p, on=key, how="outer").fillna(0)
    m["click_change"] = m.current_clicks - m.previous_clicks
    m["impression_change"] = m.current_impressions - m.previous_impressions
    m["position_change"] = m.previous_position - m.current_position
    return m


def pct(a, b):
    return "New" if not b and a else ("0%" if not b else f"{(a-b)/b:+.1%}")


def fmt(value, kind="number"):
    if kind == "percent":
        return f"{value:.2%}"
    if kind == "position":
        return f"{value:.1f}"
    return f"{value:,.0f}"


def shade(cell, fill="F2C94C"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def status_color(value):
    text = str(value).strip()
    lower = text.lower()
    if text.startswith("+") or text.startswith("↑") or lower in {"yes", "new", "positive", "improved", "up"}:
        return "008A3B", "E2F4E8"
    if text.startswith("-") or text.startswith("↓") or lower in {"no", "needs improvement", "declined", "down"}:
        return "C62828", "FCE4E4"
    return None


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = str(h)
        shade(t.rows[0].cells[i])
        for run in t.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            shown = str(value)
            if shown.startswith("+"):
                shown = "↑ " + shown
            elif shown.startswith("-"):
                shown = "↓ " + shown
            cells[i].text = shown
            signal = status_color(shown)
            if signal:
                font_color, fill = signal
                shade(cells[i], fill)
                for run in cells[i].paragraphs[0].runs:
                    run.font.color.rgb = RGBColor.from_string(font_color)
                    run.bold = True
    doc.add_paragraph()
    return t


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def save_bar(path, labels, values, title, ylabel, colors=None):
    plt.figure(figsize=(7.4, 3.8))
    bars = plt.bar(labels, values, color=colors or ["#F2C94C", "#6B7280"])
    plt.title(title, fontsize=15, fontweight="bold")
    plt.ylabel(ylabel)
    plt.grid(axis="y", alpha=0.2)
    for bar, value in zip(bars, values):
        plt.text(bar.get_x() + bar.get_width()/2, bar.get_height(), f"{value:,.0f}", ha="center", va="bottom", fontweight="bold")
    plt.tight_layout()
    plt.savefig(path, dpi=180, bbox_inches="tight")
    plt.close()


def main():
    cur_total, prev_total = gsc(), gsc(start=PREVIOUS_START, end=PREVIOUS_END)
    cur = cur_total.iloc[0].to_dict() if not cur_total.empty else {k: 0 for k in ("clicks", "impressions", "ctr", "position")}
    prev = prev_total.iloc[0].to_dict() if not prev_total.empty else {k: 0 for k in ("clicks", "impressions", "ctr", "position")}

    reports = {}
    for name, dims in {
        "queries": ["query"], "pages": ["page"], "countries": ["country"], "devices": ["device"],
        "search_appearance": ["searchAppearance"], "query_page": ["query", "page"],
    }.items():
        reports[f"{name}_aug"] = gsc(dims)
        reports[f"{name}_jul"] = gsc(dims, PREVIOUS_START, PREVIOUS_END)
    reports["queries_comparison"] = merge_months(reports["queries_aug"], reports["queries_jul"], "query")
    reports["pages_comparison"] = merge_months(reports["pages_aug"], reports["pages_jul"], "page")
    reports["countries_comparison"] = merge_months(reports["countries_aug"], reports["countries_jul"], "country")
    reports["devices_comparison"] = merge_months(reports["devices_aug"], reports["devices_jul"], "device")
    reports["daily_aug"] = gsc(["date"])
    reports["daily_jul"] = gsc(["date"], PREVIOUS_START, PREVIOUS_END)
    try:
        reports["image_aug"] = gsc(["query"], search_type="image")
        reports["image_jul"] = gsc(["query"], PREVIOUS_START, PREVIOUS_END, "image")
    except Exception:
        reports["image_aug"], reports["image_jul"] = pd.DataFrame(), pd.DataFrame()

    ga_files = [
        "ga4_period_comparisons", "ga4_channel_acquisition", "ga4_source_medium", "ga4_landing_pages", "ga4_pages",
        "ga4_events", "ga4_country", "ga4_region", "ga4_city", "ga4_devices", "ga4_browsers", "ga4_operating_systems",
        "ga4_language", "ga4_day_hour", "ga4_page_referrer", "ga4_google_ads_campaigns", "ga4_ecommerce_items",
        "ga4_organic_landing_pages",
    ]
    ga = {name: pd.read_csv(OUTDIR / f"{name}.csv") for name in ga_files}
    recovered_referrers = pd.DataFrame(recover_ga4(["pageReferrer"], ["sessions", "totalUsers", "screenPageViews", "keyEvents"], "2026-08-01", "2026-08-29"))
    recovered_sources = pd.DataFrame(recover_ga4(["sessionSourceMedium"], ["sessions", "totalUsers", "screenPageViews", "keyEvents"], "2026-08-01", "2026-08-29"))
    recovered_countries = pd.DataFrame(recover_ga4(["country"], ["sessions", "totalUsers", "screenPageViews", "keyEvents"], "2026-08-01", "2026-08-29"))
    for frame in (recovered_referrers, recovered_sources, recovered_countries):
        for col in ("sessions", "totalUsers", "screenPageViews", "keyEvents"):
            if col in frame:
                frame[col] = pd.to_numeric(frame[col], errors="coerce").fillna(0)
    reports["recovered_referrers"] = recovered_referrers
    reports["recovered_sources"] = recovered_sources
    reports["recovered_countries"] = recovered_countries

    clicks_chart = OUTDIR / "chart-monthly-clicks.png"
    impressions_chart = OUTDIR / "chart-monthly-impressions.png"
    countries_chart = OUTDIR / "chart-top-countries.png"
    save_bar(clicks_chart, ["August 1–27", "July 1–27"], [cur["clicks"], prev["clicks"]], "Google Search Clicks", "Clicks", ["#19A55A", "#6B7280"])
    save_bar(impressions_chart, ["August 1–27", "July 1–27"], [cur["impressions"], prev["impressions"]], "Google Search Appearances", "Impressions", ["#19A55A", "#6B7280"])
    top3 = recovered_countries.sort_values("sessions", ascending=False).head(3)
    save_bar(countries_chart, top3["country"].tolist(), top3["sessions"].tolist(), "Top 3 Countries by Recorded Visits", "Sessions", ["#F2C94C", "#2F80ED", "#19A55A"])

    with pd.ExcelWriter(XLSX, engine="openpyxl") as writer:
        pd.DataFrame([{**{"period": "August 1-27"}, **cur}, {**{"period": "July 1-27"}, **prev}]).to_excel(writer, sheet_name="KPI comparison", index=False)
        for name, frame in reports.items():
            frame.to_excel(writer, sheet_name=name[:31], index=False)

    doc = Document()
    apply_branding(doc, "AZ Web Corp")
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.65)
    sec.left_margin = sec.right_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10.5)
    for name, size, color in (("Title", 27, "18202A"), ("Heading 1", 18, "18202A"), ("Heading 2", 13, "A56F00")):
        doc.styles[name].font.name = "Aptos Display"
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor.from_string(color)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("AZ Web Corp\nMonthly Google Performance Report")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("August 1–27, 2026 compared with July 1–27, 2026\nGoogle Search Console + Google Analytics 4").bold = True

    doc.add_heading("KEY POINTS AT A GLANCE", 1)
    callout = table(doc, ["Status", "What immediately matters"], [
        ("Positive", "Cloudflare is purged and the correct AZ Web Corp analytics ID is publicly live."),
        ("Positive", "Several commercial web-development searches are already close to page one."),
        ("Needs improvement", "Google visibility is not yet producing enough qualified clicks."),
        ("Needs improvement", "GA4 records form activity, but lead events are not configured as conversions."),
    ])
    for row in callout.rows:
        for cell in row.cells:
            for run in cell.paragraphs[0].runs:
                run.bold = True

    doc.add_heading("Executive summary", 1)
    doc.add_paragraph(
        f"During August 1–27, AZ Web Corp received {fmt(cur['clicks'])} Google clicks from {fmt(cur['impressions'])} search appearances. "
        f"During the matching July period, it received {fmt(prev['clicks'])} clicks from {fmt(prev['impressions'])} appearances. "
        "This report uses equal-length periods so the comparison is fair."
    )
    table(doc, ["Measure", "August 1–27", "July 1–27", "Change"], [
        ("Google clicks", fmt(cur["clicks"]), fmt(prev["clicks"]), pct(cur["clicks"], prev["clicks"])),
        ("Google impressions", fmt(cur["impressions"]), fmt(prev["impressions"]), pct(cur["impressions"], prev["impressions"])),
        ("Click-through rate", fmt(cur["ctr"], "percent"), fmt(prev["ctr"], "percent"), f"{(cur['ctr']-prev['ctr'])*100:+.2f} points"),
        ("Average Google position", fmt(cur["position"], "position"), fmt(prev["position"], "position"), f"{prev['position']-cur['position']:+.1f} positions"),
    ])
    doc.add_paragraph("Two quick graphs for non-technical readers:").bold = True
    doc.add_picture(str(clicks_chart), width=Inches(6.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(str(impressions_chart), width=Inches(6.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Where the clicks and visits came from", 1)
    doc.add_paragraph(
        "Search Console reports clicks from Google Search only. It does not include direct visits, social media, referrals or advertising. "
        "GA4 identifies the broader traffic source after a visitor reaches the website. Because correct GA4 tracking began only on August 27, the source figures below cover the first three tracked days, not the entire month."
    )
    table(doc, ["Measurement system", "Source represented", "Recorded volume", "Plain-English meaning"], [
        ("Google Search Console", "Google organic search", f"{fmt(cur['clicks'])} clicks", "People clicked AZ Web Corp in unpaid Google results"),
        ("Google Analytics 4", "Direct", "33 sessions", "Typed/bookmarked visits or visits whose source could not be identified"),
        ("Google Analytics 4", "Organic Search", "3 sessions", "Visits identified as coming from an unpaid search engine result"),
        ("Google Analytics 4", "Paid advertising", "0 sessions", "No Google Ads traffic was returned"),
        ("Google Analytics 4", "Ecommerce", "0 transactions", "No online purchase activity was returned"),
    ])
    source_medium = ga["ga4_source_medium"].sort_values("sessions", ascending=False)
    doc.add_heading("GA4 source and medium detail", 2)
    table(doc, ["Source / medium", "Sessions", "Users", "Engagement rate", "Key events"], [
        (r.sessionSourceMedium, fmt(r.sessions), fmt(r.totalUsers), fmt(r.engagementRate, "percent"), fmt(r.keyEvents)) for _, r in source_medium.iterrows()
    ])
    referrers = ga["ga4_page_referrer"].sort_values("sessions", ascending=False).head(10)
    doc.add_heading("Recorded referring pages", 2)
    table(doc, ["Referring page", "Sessions", "Users", "Page views", "Key events"], [
        (r.pageReferrer or "Direct / not provided", fmt(r.sessions), fmt(r.totalUsers), fmt(r.screenPageViews), fmt(r.keyEvents)) for _, r in referrers.iterrows()
    ])
    doc.add_paragraph(
        "Important: ‘Direct’ does not always mean the visitor manually typed the address. It also includes traffic where campaign tags or referral information were unavailable."
    )
    doc.add_heading("Partner and referral links", 2)
    table(doc, ["Referral source", "Link status", "GA4-attributed sessions", "Explanation"], [
        ("Everything IT (everythingit.ie)", "Positive", "Not separately attributed", "Verified live footer link to AZ Web Corp on the homepage and location pages"),
        ("Google", "Positive", "2 referring-page sessions", "Google appears as a recorded external referring page in recovered data"),
        ("Bing", "Positive", "1 referring-page session", "Bing appears as a recorded external referring page in recovered data"),
    ])
    doc.add_paragraph(
        "Why Everything IT has no exact GA4 count: the footer backlink is active and capable of passing referral information, but the available Analytics rows do not name everythingit.ie. During most of the period, AZ Web Corp was sending analytics to the wrong property, and many recovered visits are classified as Unassigned or Direct. The report therefore confirms the backlink without claiming an unsupported number."
    )

    qcmp = reports["queries_comparison"]
    top_aug = reports["queries_aug"].sort_values(["clicks", "impressions"], ascending=False).head(15)
    doc.add_heading("1. Search queries", 1)
    doc.add_paragraph("These are the searches that produced the most clicks and visibility during August.")
    table(doc, ["Search query", "Clicks", "Impressions", "CTR", "Position"], [
        (r["query"], fmt(r["clicks"]), fmt(r["impressions"]), fmt(r["ctr"], "percent"), fmt(r["position"], "position")) for _, r in top_aug.iterrows()
    ])
    winners = qcmp.sort_values(["click_change", "impression_change"], ascending=False).head(10)
    losses = qcmp.sort_values(["click_change", "impression_change"], ascending=True).head(10)
    doc.add_heading("Queries gaining month over month", 2)
    table(doc, ["Query", "Click change", "Impression change", "Position improvement"], [
        (r["query"], f"{r['click_change']:+.0f}", f"{r['impression_change']:+.0f}", f"{r['position_change']:+.1f}") for _, r in winners.iterrows()
    ])
    doc.add_heading("Queries losing month over month", 2)
    table(doc, ["Query", "Click change", "Impression change", "Position improvement"], [
        (r["query"], f"{r['click_change']:+.0f}", f"{r['impression_change']:+.0f}", f"{r['position_change']:+.1f}") for _, r in losses.iterrows()
    ])

    pcmp = reports["pages_comparison"]
    doc.add_heading("2. Landing pages from Google", 1)
    doc.add_paragraph("This section shows which pages gained or lost Google visibility and clicks between the two months.")
    page_rows = pcmp.sort_values("current_impressions", ascending=False).head(20)
    table(doc, ["Page", "Aug clicks", "Jul clicks", "Click change", "Aug impressions", "Jul impressions", "Impression change"], [
        (r["page"].replace("https://azwebcorp.com", "") or "/", fmt(r["current_clicks"]), fmt(r["previous_clicks"]), f"{r['click_change']:+.0f}", fmt(r["current_impressions"]), fmt(r["previous_impressions"]), f"{r['impression_change']:+.0f}") for _, r in page_rows.iterrows()
    ])

    doc.add_heading("3. Countries and devices", 1)
    countries = reports["countries_comparison"].sort_values("current_impressions", ascending=False).head(15)
    table(doc, ["Country", "Aug clicks", "Jul clicks", "Click change", "Aug impressions", "Jul impressions", "Impression change"], [
        (r["country"], fmt(r["current_clicks"]), fmt(r["previous_clicks"]), f"{r['click_change']:+.0f}", fmt(r["current_impressions"]), fmt(r["previous_impressions"]), f"{r['impression_change']:+.0f}") for _, r in countries.iterrows()
    ])
    devices = reports["devices_comparison"].sort_values("current_impressions", ascending=False)
    table(doc, ["Device", "Aug clicks", "Jul clicks", "Click change", "Aug impressions", "Jul impressions", "Impression change", "Aug CTR"], [
        (r["device"].title(), fmt(r["current_clicks"]), fmt(r["previous_clicks"]), f"{r['click_change']:+.0f}", fmt(r["current_impressions"]), fmt(r["previous_impressions"]), f"{r['impression_change']:+.0f}", fmt(r["current_ctr"], "percent")) for _, r in devices.iterrows()
    ])
    doc.add_heading("Top three countries by recorded website traffic", 2)
    doc.add_paragraph("These GA4 sessions were recovered by filtering the previously incorrect property specifically for the azwebcorp.com hostname.")
    table(doc, ["Rank", "Country", "Sessions", "Users", "Page views"], [
        (i + 1, r.country, fmt(r.sessions), fmt(r.totalUsers), fmt(r.screenPageViews)) for i, (_, r) in enumerate(top3.iterrows())
    ])
    doc.add_picture(str(countries_chart), width=Inches(6.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("4. Search appearance and image search", 1)
    app_rows = len(reports["search_appearance_aug"])
    image_impressions = reports["image_aug"]["impressions"].sum() if not reports["image_aug"].empty else 0
    doc.add_paragraph(
        f"Google returned {app_rows} special search-appearance categories for August. Image Search generated {fmt(image_impressions)} query-level impressions in the export. "
        "The supporting workbook contains the complete image-query tables for both months."
    )

    doc.add_heading("5. Google Analytics monthly comparison", 1)
    doc.add_paragraph(
        "The correct AZ Web Corp GA4 property began receiving data on August 27. Therefore, July contains no GA4 data and August contains only three days. "
        "The figures below are a tracking baseline, not a complete month-over-month business comparison."
    )
    comp = ga["ga4_period_comparisons"].query("window_days == 28")
    current_ga = comp.query("period == 'current'").iloc[0]
    table(doc, ["GA4 measure", "August recorded", "July recorded", "Meaning"], [
        ("Sessions", fmt(current_ga.sessions), "0", "36 visits were tracked after collection began"),
        ("Users", fmt(current_ga.totalUsers), "0", "28 people were recorded"),
        ("Page views", fmt(current_ga.screenPageViews), "0", "58 pages were viewed"),
        ("Engagement rate", fmt(current_ga.engagementRate, "percent"), "0%", "44.44% of sessions were engaged"),
        ("Key events / conversions", fmt(current_ga.keyEvents), "0", "Conversions are not configured yet"),
    ])

    doc.add_heading("6. Acquisition: how visitors arrived", 1)
    acq = ga["ga4_channel_acquisition"].sort_values("sessions", ascending=False)
    table(doc, ["Channel", "Sessions", "Users", "Engagement rate", "Key events"], [
        (r.sessionDefaultChannelGroup, fmt(r.sessions), fmt(r.totalUsers), fmt(r.engagementRate, "percent"), fmt(r.keyEvents)) for _, r in acq.iterrows()
    ])

    doc.add_heading("7. GA4 landing pages and content", 1)
    landing = ga["ga4_landing_pages"].sort_values("sessions", ascending=False).head(15)
    table(doc, ["Landing page", "Sessions", "Users", "Page views", "Engagement"], [
        (r.landingPagePlusQueryString, fmt(r.sessions), fmt(r.totalUsers), fmt(r.screenPageViews), fmt(r.engagementRate, "percent")) for _, r in landing.iterrows()
    ])

    doc.add_heading("8. Events and lead activity", 1)
    events = ga["ga4_events"].sort_values("eventCount", ascending=False)
    table(doc, ["Event", "Times recorded", "Users", "Marked as conversion?"], [
        (r.eventName, fmt(r.eventCount), fmt(r.totalUsers), "Yes" if r.keyEvents else "No") for _, r in events.iterrows()
    ])
    doc.add_paragraph("Important: GA4 recorded 14 form starts and one form submission, but form submission is not marked as a key event. Reported conversions therefore remain zero.")

    doc.add_heading("9. Audience location and technology", 1)
    doc.add_paragraph("The workbook contains full country, region, city, device, browser, operating system, language, screen-resolution and hour-of-day tables. Early GA4 counts are too small for reliable strategic conclusions.")
    table(doc, ["Report category", "Rows available"], [
        ("Countries", len(ga["ga4_country"])), ("Regions", len(ga["ga4_region"])), ("Cities", len(ga["ga4_city"])),
        ("Devices", len(ga["ga4_devices"])), ("Browsers", len(ga["ga4_browsers"])), ("Operating systems", len(ga["ga4_operating_systems"])),
        ("Languages", len(ga["ga4_language"])), ("Day and hour combinations", len(ga["ga4_day_hour"])),
    ])

    doc.add_heading("10. Advertising, ecommerce and referrals", 1)
    bullets(doc, [
        f"Google Ads campaign rows: {len(ga['ga4_google_ads_campaigns'])}. No linked advertising activity was returned.",
        f"Ecommerce item rows: {len(ga['ga4_ecommerce_items'])}. No ecommerce transactions were returned.",
        f"Referrer rows: {len(ga['ga4_page_referrer'])}. The detailed sources are preserved in the supporting workbook.",
        f"Organic landing-page rows: {len(ga['ga4_organic_landing_pages'])}. Organic GA4 history is currently limited to the new three-day baseline.",
    ])

    doc.add_heading("11. Actions recommended from this monthly comparison", 1)
    table(doc, ["Priority", "Action", "Why"], [
        ("Completed", "Cloudflare purged and G-R5RNDSH327 verified publicly", "The correct AZ Web Corp stream is now served to visitors and the Lootertech ID is absent"),
        ("Immediate", "Mark form_submit, phone clicks and email clicks as GA4 key events", "Monthly reports must show leads, not only visits"),
        ("This week", "Improve pages and snippets for queries ranking positions 4–20", "These are the fastest opportunities for qualified clicks"),
        ("This week", "Review pages that lost impressions or clicks from July to August", "Stops avoidable month-over-month decline"),
        ("Monthly", "Repeat equal-day monthly comparisons until complete months are available", "Keeps reporting fair while Search Console data is delayed"),
    ])

    doc.add_heading("12. Report inventory", 1)
    doc.add_paragraph(
        "This master report covers Google Search performance, queries, pages, query-to-page relationships, countries, devices, search appearance, image search, GA4 acquisition, campaigns, landing pages, content, events, conversions, geography, technology, referrals, Google Ads and ecommerce. "
        "The accompanying monthly Excel workbook retains the exact underlying rows for audit and filtering."
    )
    doc.add_heading("Tracking ownership", 2)
    table(doc, ["Item", "Confirmed value"], [
        ("Google account", "yasirmemon1976@gmail.com"), ("Access level", "Administrator"),
        ("GA4 property", "247570709"), ("Correct stream", "G-R5RNDSH327"),
    ])
    doc.add_heading("Overall assessment", 1)
    doc.add_paragraph(
        "AZ Web Corp has meaningful search visibility and several realistic page-one opportunities, but the site is not yet converting that visibility into enough qualified clicks or measurable leads. The month should be viewed as a rebuilding and measurement-baseline period rather than a finished success story."
    )
    assessment_rows = [
        ("Search visibility", "Positive" if cur["impressions"] >= prev["impressions"] else "Needs improvement", "Google appearances are " + ("higher" if cur["impressions"] >= prev["impressions"] else "lower") + " than the matching July period."),
        ("Search clicks", "Positive" if cur["clicks"] >= prev["clicks"] else "Needs improvement", "Clicks are " + ("higher" if cur["clicks"] >= prev["clicks"] else "lower") + " than the matching July period."),
        ("Referral visibility", "Positive", "Everything IT's footer backlink is verified live, alongside recorded Google and Bing referrals."),
        ("Analytics correction", "Positive", "Cloudflare was purged; the correct AZ Web Corp ID is publicly live and the Lootertech ID is absent. Earlier-period data still carries the documented historical limitation."),
        ("Lead measurement", "Needs improvement", "Form activity exists, but key events/conversions are not configured."),
        ("Growth opportunity", "Positive", "Several commercial web-development searches already rank close to page one."),
    ]
    table(doc, ["Area", "Assessment", "What it means"], assessment_rows)
    doc.add_paragraph(
        "Overall rating: PROMISING, BUT TRACKING AND CONVERSION MEASUREMENT NEED IMMEDIATE ATTENTION.",
    ).runs[0].font.color.rgb = RGBColor.from_string("C58A00")
    doc.add_heading("What the customer should remember", 2)
    bullets(doc, [
        "AZ Web Corp is becoming more visible in Google, with several valuable searches close to page one.",
        "More of that visibility needs to turn into qualified website visits and completed enquiries.",
        "Analytics is now pointed to the correct account, but conversions must be configured before lead performance can be judged accurately.",
        "The next month should focus on near-page-one web-development keywords, stronger Google snippets and reliable lead tracking.",
    ])
    doc.save(DOCX)
    print(DOCX)
    print(XLSX)


if __name__ == "__main__":
    main()
