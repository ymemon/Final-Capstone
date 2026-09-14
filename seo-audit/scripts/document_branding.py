from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

LOGO = Path(r"C:\Users\yasir\Documents\Final-Capstone\seo-audit\assets\branding\Azwebcorp-black_logo.png")
CONTACT = "AZ Web Corp  |  (480) 818-5761  |  info@azwebcorp.com  |  azwebcorp.com"


def add_page_field(paragraph):
    paragraph.add_run("  |  Page ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText"); text.set(qn("xml:space"), "preserve"); text.text = "PAGE"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, end])


def add_watermark(header, text="AZ WEB CORP"):
    xml = f'''<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
      xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">
      <w:pPr><w:jc w:val="center"/></w:pPr>
      <w:r><w:pict>
        <v:shape id="PowerPlusWaterMarkObject357476642" o:spid="_x0000_s1025"
          type="#_x0000_t136"
          style="position:absolute;margin-left:0;margin-top:0;width:500pt;height:120pt;rotation:315;z-index:-251654144;mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin"
          fillcolor="#E8E8E8" stroked="f" o:allowincell="f">
          <v:textpath style="font-family:'Aptos Display';font-size:1pt;font-weight:bold" string="{text}"/>
        </v:shape>
      </w:pict></w:r>
    </w:p>'''
    header._element.append(parse_xml(xml))


def apply_branding(doc, report_for):
    if not LOGO.exists():
        raise FileNotFoundError(LOGO)

    logo = doc.add_picture(str(LOGO), width=Inches(3.15))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = label.add_run(f"Prepared by AZ Web Corp for {report_for}")
    run.bold = True; run.font.size = Pt(9); run.font.color.rgb = RGBColor(128, 92, 0)

    for section in doc.sections:
        section.header_distance = Inches(0.25)
        section.footer_distance = Inches(0.3)
        add_watermark(section.header)
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(CONTACT)
        r.bold = True; r.font.size = Pt(8); r.font.color.rgb = RGBColor(90, 90, 90)
        add_page_field(p)
        for run in p.runs:
            run.font.size = Pt(8); run.font.color.rgb = RGBColor(90, 90, 90)
