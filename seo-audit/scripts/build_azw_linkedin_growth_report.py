from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from document_branding import apply_branding

OUT = Path(r"C:\Users\yasir\Documents\Final-Capstone\seo-audit\reports\AZWebCorp-LinkedIn-Growth-2026-08-30")
DOCX = OUT / "AZ-Web-Corp-GA4-Growth-LinkedIn-Report.docx"
CAPTION = OUT / "LinkedIn-Post-Caption.txt"

GREEN, GOLD, DARK, RED, GRAY = "#19A55A", "#F2C94C", "#18202A", "#C62828", "#6B7280"


def chart_compare(path, labels, previous, current, title):
    x = np.arange(len(labels)); width = .36
    fig, ax = plt.subplots(figsize=(8.6, 4.8))
    b1 = ax.bar(x-width/2, previous, width, label="Aug 16–22", color=GRAY)
    b2 = ax.bar(x+width/2, current, width, label="Aug 23–29", color=GREEN)
    ax.set_title(title, fontsize=17, fontweight="bold", color=DARK)
    ax.set_xticks(x, labels); ax.legend(frameon=False); ax.grid(axis="y", alpha=.18)
    for bars in (b1,b2):
        for b in bars: ax.text(b.get_x()+b.get_width()/2,b.get_height(),f"{b.get_height():,.0f}",ha="center",va="bottom",fontweight="bold")
    fig.tight_layout(); fig.savefig(path,dpi=190,bbox_inches="tight"); plt.close(fig)


def chart_growth(path):
    labels=["Organic Search","Referral","Page views","Events","Active users","Sessions","Direct","Views/user"]
    values=[600,440,237,209.7,172,171,111.8,23.9]
    fig,ax=plt.subplots(figsize=(8.6,5.3)); bars=ax.barh(labels[::-1],values[::-1],color=GREEN)
    ax.set_title("Week-over-Week Growth",fontsize=17,fontweight="bold",color=DARK); ax.set_xlabel("Percentage increase"); ax.grid(axis="x",alpha=.18)
    for b,v in zip(bars,values[::-1]): ax.text(b.get_width()+8,b.get_y()+b.get_height()/2,f"↑ {v:g}%",va="center",fontweight="bold",color="#087A36")
    fig.tight_layout(); fig.savefig(path,dpi=190,bbox_inches="tight"); plt.close(fig)


def chart_sources(path):
    labels=["Direct","Google organic","EverythingIT.ie","Ahrefs","Other / unidentified","Bing","Facebook","LinkedIn","ChatGPT / AI"]
    values=[161,27,22,11,5,3,2,2,1]
    fig,ax=plt.subplots(figsize=(8.6,5.4)); colors=[GOLD,GREEN,"#2F80ED","#8B5CF6",GRAY,"#4C8BF5","#1877F2","#0A66C2","#10A37F"]
    bars=ax.barh(labels[::-1],values[::-1],color=colors[::-1]); ax.set_title("Where 234 Monthly Sessions Came From",fontsize=17,fontweight="bold",color=DARK); ax.set_xlabel("Sessions"); ax.grid(axis="x",alpha=.18)
    for b,v in zip(bars,values[::-1]): ax.text(b.get_width()+1,b.get_y()+b.get_height()/2,str(v),va="center",fontweight="bold")
    fig.tight_layout(); fig.savefig(path,dpi=190,bbox_inches="tight"); plt.close(fig)


def chart_countries(path):
    labels=["United States","Ireland","Netherlands","Pakistan"]; values=[12,6,5,4]
    fig,ax=plt.subplots(figsize=(8.6,4.7)); bars=ax.bar(labels,values,color=[GREEN,"#2F80ED","#F2994A","#8B5CF6"])
    ax.set_title("Increase in Active Users by Country",fontsize=17,fontweight="bold",color=DARK); ax.set_ylabel("Additional active users"); ax.grid(axis="y",alpha=.18)
    for b,v in zip(bars,values): ax.text(b.get_x()+b.get_width()/2,b.get_height(),f"↑ {v}",ha="center",va="bottom",fontweight="bold",color="#087A36")
    fig.tight_layout(); fig.savefig(path,dpi=190,bbox_inches="tight"); plt.close(fig)


def add_picture(doc,path):
    doc.add_picture(str(path),width=Inches(6.75)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER


def heading_box(doc,title,text,color="19A55A"):
    t=doc.add_table(rows=1,cols=1); t.style="Table Grid"; cell=t.cell(0,0); cell.text=""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),"E2F4E8" if color=="19A55A" else "FFF4D6"); pr.append(shd)
    p=cell.paragraphs[0]; r=p.add_run(title+"\n"); r.bold=True; r.font.size=Pt(15); r.font.color.rgb=RGBColor.from_string(color)
    r=p.add_run(text); r.bold=True; r.font.size=Pt(11)
    doc.add_paragraph()


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    traffic=OUT/"weekly-traffic.png"; engagement=OUT/"weekly-engagement.png"; growth=OUT/"weekly-growth.png"; sources=OUT/"monthly-sources.png"; countries=OUT/"country-growth.png"
    chart_compare(traffic,["Active users","Sessions"],[25,42],[68,114],"More People Reached the Website")
    chart_compare(engagement,["Page views","Events"],[100,288],[337,892],"Visitors Did More on the Website")
    chart_growth(growth); chart_sources(sources); chart_countries(countries)

    doc=Document(); apply_branding(doc,"AZ Web Corp LinkedIn audience"); sec=doc.sections[0]; sec.top_margin=sec.bottom_margin=Inches(.6); sec.left_margin=sec.right_margin=Inches(.7)
    doc.styles["Normal"].font.name="Aptos"; doc.styles["Normal"].font.size=Pt(11)
    for name,size,color in (("Title",27,"18202A"),("Heading 1",19,"18202A"),("Heading 2",14,"A56F00")):
        doc.styles[name].font.name="Aptos Display"; doc.styles[name].font.size=Pt(size); doc.styles[name].font.color.rgb=RGBColor.from_string(color)
    p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("AZ Web Corp\nDigital Growth Snapshot")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("GA4 results: August 23–29 vs. August 16–22\nFull-month acquisition: July 31–August 29").bold=True
    heading_box(doc,"↑ ORGANIC SEARCH SESSIONS GREW 600%","Unpaid search traffic increased from 2 to 14 sessions in one week.")
    heading_box(doc,"↑ REFERRAL SESSIONS GREW 440%","Referral traffic increased from 5 to 27 sessions. EverythingIT.ie generated 22 sessions during the full-month period.")
    heading_box(doc,"TRACKING NOTE","Percentages are large because the starting numbers were small. The direction is encouraging, but sustained growth over several weeks matters more than one isolated spike.","A56F00")

    doc.add_heading("The week in one view",1)
    add_picture(doc,traffic); add_picture(doc,engagement); add_picture(doc,growth)
    doc.add_heading("What improved",1)
    rows=[("Active users","25","68","↑ +43 (+172%)"),("Sessions","42","114","↑ +72 (+171%)"),("Page views","100","337","↑ +237 (+237%)"),("Views per user","4.00","4.96","↑ +0.96 (+23.9%)"),("Events","288","892","↑ +604 (+209.7%)"),("Organic Search sessions","2","14","↑ +12 (+600%)"),("Referral sessions","5","27","↑ +22 (+440%)"),("Direct sessions","34","72","↑ +38 (+111.8%)")]
    t=doc.add_table(rows=1,cols=4); t.style="Table Grid"
    for i,h in enumerate(["Metric","Previous week","Current week","Change"]): t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=v
        for run in cells[3].paragraphs[0].runs: run.bold=True; run.font.color.rgb=RGBColor.from_string("008A3B")

    doc.add_heading("Where the full month's traffic came from",1)
    doc.add_paragraph("The monthly view recorded 234 sessions, up from 79—an increase of 155 sessions or 196%.")
    add_picture(doc,sources)
    doc.add_heading("The referral win",2)
    doc.add_paragraph("EverythingIT.ie delivered 22 referral sessions. The most likely path is the AZ Web Corp footer credit link displayed across Everything IT's website and location pages. This is measurable evidence that a live partner backlink can generate website visits—not merely link authority.")
    doc.add_paragraph("Ahrefs generated 11 sessions. These may include SEO research or team activity, so they should not automatically be counted as prospective customers.")
    doc.add_paragraph("Direct accounted for 161 sessions. GA4 uses Direct when it cannot identify another source, so this should not all be interpreted as brand awareness.")

    doc.add_heading("Search traffic is moving",1)
    heading_box(doc,"↑ GOOGLE ORGANIC: 13 → 27 MONTHLY SESSIONS","That is an increase of 14 sessions, or approximately 107.7%, across the full monthly comparison.")
    doc.add_paragraph("The weekly organic increase of 2 to 14 sessions supports the same direction. This is encouraging SEO evidence, but the campaign should be judged on whether organic sessions continue rising over the next several weeks.")

    doc.add_heading("Geographic growth",1); add_picture(doc,countries)
    doc.add_paragraph("U.S. growth is the most commercially relevant for AZ Web Corp. Ireland may partly reflect partner activity involving Everything IT. Pakistan may include developer or team activity. Netherlands traffic should be monitored for quality, bots or VPN usage before treating it as prospective business.")

    doc.add_heading("Overall assessment",1)
    heading_box(doc,"PROMISING AND MEASURABLE GROWTH","Users, sessions, page views, engagement, organic search and referrals all improved in the same direction. That is a stronger signal than a single-day spike.")
    doc.add_paragraph("The results are encouraging, not conclusive. Organic volume is still modest, Direct attribution is unusually high, and some referral activity may come from tools or team testing. The next priority is to sustain organic growth, identify every referral source, and measure completed leads—not only visits and events.")
    doc.add_heading("What to watch next",2)
    for text in ["Whether Organic Search continues from 14 toward 20, 35 and 50+ weekly sessions.","Which Google queries and landing pages generate the organic visits.","Whether Everything IT referrals engage and convert into enquiries.","Completed forms, phone clicks and email clicks configured as GA4 key events.","U.S. traffic growth separated from team, bot, VPN and tool activity."]: doc.add_paragraph(text,style="List Bullet")
    doc.add_paragraph("Source: Google Analytics 4 figures supplied by AZ Web Corp. Charts were recreated from the supplied GA4 figures for clear LinkedIn presentation.").italic=True
    doc.save(DOCX)

    CAPTION.write_text("""SEO progress is not one flashy percentage. It is several useful signals moving together.\n\nDuring August 23–29, AZ Web Corp recorded:\n\n• 68 active users — up 172%\n• 114 sessions — up 171%\n• 337 page views — up 237%\n• 892 events — up 209.7%\n• 14 Organic Search sessions — up 600%\n• 27 referral sessions — up 440%\n\nAcross the full monthly view, Google Organic increased from 13 to 27 sessions, while EverythingIT.ie generated 22 referral sessions through our live partner link.\n\nThe starting numbers are still modest, so this is not a victory lap. The real test is whether qualified organic traffic continues growing—and whether those visits become enquiries. But users, engagement, search and referrals all improving together is a promising direction.\n\n#SEO #GoogleAnalytics #OrganicGrowth #DigitalMarketing #WebDevelopment #ArizonaBusiness #AZWebCorp\n""",encoding="utf-8")
    print(DOCX); print(CAPTION)


if __name__=="__main__": main()
