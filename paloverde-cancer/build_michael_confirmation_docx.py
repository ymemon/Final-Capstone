from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"C:\Users\yasir\Documents\Final-Capstone")
PV = ROOT / "paloverde-cancer"
SHOTS = PV / "screenshots"
OUT = PV / "Palo-Verde-Website-Corrections-Confirmation-2026-08-31.docx"
LOGO = ROOT / "seo-audit" / "content" / "Azwebcorp-light_logo.png"

URLS = [
    ("Your Team", "https://875051.us16.myftpupload.com/your-team/"),
    ("Dr. Mamani", "https://875051.us16.myftpupload.com/your-team/dr-mamani/"),
    ("Estrella", "https://875051.us16.myftpupload.com/estrella-location/"),
    ("Glendale", "https://875051.us16.myftpupload.com/glendale-location/"),
    ("Scottsdale", "https://875051.us16.myftpupload.com/scottsdale-location/"),
    ("Gilbert / East Valley", "https://875051.us16.myftpupload.com/east-valley-location/"),
    ("PET Scan Imaging", "https://875051.us16.myftpupload.com/pet-scan-imaging/"),
]

IMAGES = [
    ("Updated Your Team page — Dr. Mamani is included", "michael-confirmation-team-mamani.png"),
    ("Gilbert office added to the homepage photo rotation", "michael-confirmation-gilbert-carousel.png"),
    ("Estrella full-resolution office photo", "michael-confirmation-estrella-office-photo.png"),
    ("Glendale full-resolution office photo", "michael-confirmation-glendale-office-photo.png"),
    ("Scottsdale full-resolution office photo", "michael-confirmation-scottsdale-office-photo.png"),
    ("Gilbert / East Valley full-resolution office photo", "michael-confirmation-gilbert-office-photo.png"),
    ("Rebuilt PET Scan Imaging page", "michael-confirmation-pet-page.png"),
    ("Dedicated PET Scan Imaging map", "michael-confirmation-pet-map.png"),
]


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(color)
    properties.append(underline)
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.6)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(24)
styles["Title"].font.color.rgb = RGBColor(15, 42, 66)
styles["Heading 1"].font.color.rgb = RGBColor(15, 42, 66)
styles["Heading 2"].font.color.rgb = RGBColor(15, 42, 66)

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
if LOGO.exists():
    hp.add_run().add_picture(str(LOGO), width=Inches(2.2))
watermark = header.add_paragraph("AZ WEB CORP")
watermark.alignment = WD_ALIGN_PARAGRAPH.RIGHT
watermark_run = watermark.runs[0]
watermark_run.font.size = Pt(8)
watermark_run.font.color.rgb = RGBColor(205, 210, 218)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run("AZ Web Corp  |  (480) 818-5761  |  info@azwebcorp.com  |  azwebcorp.com")
fp.runs[0].font.size = Pt(8)
fp.runs[0].font.color.rgb = RGBColor(90, 98, 110)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Palo Verde Website Corrections — Confirmation")

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Prepared for Michael Bustard, IT Director • August 31, 2026").italic = True

doc.add_heading("Email draft", level=1)
doc.add_paragraph("To: Michael Bustard <mbustard@pvcancer.com>")
doc.add_paragraph("Subject: Palo Verde website corrections completed — ready for your review")
doc.add_paragraph("Hi Mike,")
doc.add_paragraph("The corrections from your list have been completed on the Palo Verde website.")

items = [
    "Confirmed the physician assignments: WVO / Estrella — Dr. Zafar, Dr. Mamani and Dr. Rakkar; TBO / Glendale — Dr. Rakkar, Dr. Mamani, Dr. Grover and Dr. Ahmad; SDO / Scottsdale — Dr. Halepota, Dr. Grover and Dr. Ahmad; GTO / Gilbert — Dr. Grover and Dr. Halepota.",
    "Added Dr. Demetrio Mamani to the Your Team page and confirmed that his profile opens correctly.",
    "Added the Gilbert office to the homepage location-photo rotation.",
    "Replaced the old low-resolution office photos for Estrella, Glendale, Scottsdale and Gilbert with full-resolution images.",
    "Rebuilt the PET Scan Imaging page with a consistent location-detail layout, the PET address and a dedicated map. The unrelated list of every office was removed from the bottom.",
    "Checked the updated pages on desktop and mobile after publishing.",
]
for item in items:
    doc.add_paragraph(item, style="List Bullet")

doc.add_paragraph("Please review the pages using the clickable links below and let me know if anything else needs to be adjusted.")
doc.add_paragraph("Thanks,\nYasir\nAZ Web Corp")

doc.add_heading("Clickable review links", level=1)
for label, url in URLS:
    p = doc.add_paragraph(style="List Bullet")
    add_hyperlink(p, f"{label}: {url}", url)

doc.add_page_break()
doc.add_heading("Confirmation screenshots", level=1)
for index, (caption, filename) in enumerate(IMAGES):
    image = SHOTS / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(caption)
    run.bold = True
    run.font.color.rgb = RGBColor(15, 42, 66)
    if image.exists():
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(str(image), width=Inches(6.7))
    if index != len(IMAGES) - 1:
        doc.add_paragraph()

doc.save(OUT)
print(OUT)
